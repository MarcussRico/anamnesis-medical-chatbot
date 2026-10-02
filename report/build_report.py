"""Builds the PBL report (.docx) in the Chennai Institute of Technology format.

    python report/build_report.py [pages.json]

pages.json (optional) maps TOC/figure/table keys to printed page numbers; it
is produced by report/paginate.py from the rendered PDF.
"""
from __future__ import annotations

import copy
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "figures"
ASSET = ROOT / "report" / "assets"
R = json.loads((ROOT / "results" / "results.json").read_text())
PAGES = json.loads(Path(sys.argv[1]).read_text()) if len(sys.argv) > 1 and Path(sys.argv[1]).exists() else {}

TITLE = "ANAMNESIS – An Information-Gain Driven Medical Chatbot for Symptom-Based Disease Triage"
TITLE_SHORT = "ANAMNESIS – An Information-Gain Driven Medical Chatbot for Symptom-Based Disease Triage"
M1, R1 = "SAI CHARAN A", "2104251040844"
M2, R2 = "ASHFAQ MUHAMMED A", "2104251040096"
m1, m2 = "Sai Charan A", "Ashfaq Muhammed A"
REPO = "https://github.com/MarcussRico/anamnesis-medical-chatbot"
FONT = "Times New Roman"
ORANGE = "ED7D31"


def pct(x, d=1):
    return f"{x * 100:.{d}f}%"


E1, E2, E3, E4 = R["e1"], R["e2"], R["e3"], R["e4"]
ES = E3["early_stop"]
NZ = R["e3_noise"]
CAL = E3["calibration"]
AB = {a["tau"]: a for a in CAL["abstention"]}
NLU_F = E4["+ fuzzy char n-gram (final)"]
NLU_L = E4["+ lay-term lexicon"]
NLU_E = E4["Exact canonical names only"]

doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.space_before = Pt(0)

sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = Cm(3.0)
sec.right_margin = Cm(2.2)
sec.top_margin = Cm(2.4)
sec.bottom_margin = Cm(2.2)
sec.header_distance = Cm(0.8)


# ---------------------------------------------------------------- primitives
def para(text="", size=12, bold=False, italic=False, align="justify", before=0, after=6, line=1.5, color=None,
         keep=False):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    if keep:
        pf.keep_with_next = True
    if text:
        rich(p, text, size, bold, italic, color)
    return p


def rich(p, text, size=12, bold=False, italic=False, color=None):
    """**bold** spans inside text are honoured."""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        r.font.size = Pt(size)
        r.font.name = FONT
        r.bold = bold or (i % 2 == 1)
        r.italic = italic
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    return p


def bullet(text, size=12):
    p = para("", align="justify", after=4, line=1.5)
    pf = p.paragraph_format
    pf.left_indent, pf.first_line_indent = Cm(0.9), Cm(-0.5)
    rich(p, "• " + text, size)
    return p


def heading_chapter(num, title, key):
    para(f"CHAPTER {num}", 16, True, align="center", after=6, keep=True)
    p = para(title.upper(), 16, True, align="center", after=14, keep=True)
    bookmark(key)
    return p


def h2(text, key=None):
    para(text, 14, True, align="left", before=10, after=6, keep=True)
    if key:
        bookmark(key)


def h3(text):
    para(text, 12.5, True, italic=True, align="left", before=6, after=4, keep=True)


def bookmark(key):
    # pagination markers are found by text in the PDF, so nothing to do here
    return key


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def picture(path, width_cm, caption=None, key=None):
    p = para("", align="center", after=2, line=1.0, keep=True)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    if caption:
        para(caption, 11, True, italic=True, align="center", after=10, line=1.15)


def caption(text):
    para(text, 11, True, italic=True, align="center", before=4, after=10, line=1.15)


def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def table(headers, rows, widths_cm, size=10.5, header_bg="D9D9D9", align_center_cols=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        rich(c.paragraphs[0], h, size, True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_bg(c, header_bg)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            rich(cells[i].paragraphs[0], str(v), size)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i in align_center_cols else WD_ALIGN_PARAGRAPH.LEFT
            cells[i].paragraphs[0].paragraph_format.line_spacing = 1.1
    for row in t.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
    fix_widths(t, widths_cm)
    # repeat header row
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    trPr.append(th)
    return t


def fix_widths(t, widths_cm):
    t.autofit = False
    tbl = t._tbl
    tblPr = tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for el in tblPr.findall(qn(tag)):
            tblPr.remove(el)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tw.set(qn("w:type"), "dxa")
    tblPr.append(tw)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    for gc, w in zip(tbl.tblGrid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for row in t.rows:
        for c, w in zip(row.cells, widths_cm):
            c.width = Cm(w)


def code(text, size=8.5):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    fix_widths(t, [15.6])
    set_cell_bg(c, "F2F2F2")
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    for i, line in enumerate(text.strip("\n").split("\n")):
        r = p.add_run(line)
        r.font.name = "Courier New"
        r.font.size = Pt(size)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        if i < len(text.strip("\n").split("\n")) - 1:
            r.add_break()
    return t


def orange_box(paragraphs, size=13):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    fix_widths(t, [15.6])
    tcPr = c._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "double")
        el.set(qn("w:sz"), "18")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), ORANGE)
        borders.append(el)
    tcPr.append(borders)
    mar = OxmlElement("w:tcMar")
    for edge, v in (("top", 160), ("bottom", 160), ("left", 220), ("right", 220)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)
    c.text = ""
    first = True
    for label, text in paragraphs:
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.6
        p.paragraph_format.space_after = Pt(4)
        if label:
            p.paragraph_format.left_indent = Cm(1.25)
            p.paragraph_format.first_line_indent = Cm(-1.25)
            r = p.add_run(label + ". ")
            r.bold = True
            r.font.size = Pt(size)
            r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00) if label.startswith("IM") else None
        r = p.add_run(text)
        r.font.size = Pt(size)


def toc_line(text, key, bold=False, indent=0.0, size=12):
    p = para("", align="left", after=3, line=1.15)
    pf = p.paragraph_format
    pf.left_indent = Cm(indent)
    pf.tab_stops.add_tab_stop(Cm(15.7), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    rich(p, text, size, bold)
    r = p.add_run("\t" + str(PAGES.get(key, "00")))
    r.font.size = Pt(size)
    r.bold = bold


# ------------------------------------------------- section/header helpers
def new_section(start_type=WD_SECTION.NEW_PAGE):
    s = doc.add_section(start_type)
    return s


def unlink(section):
    for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
        part.is_linked_to_previous = False


def clear_header(section):
    for part in (section.header, section.footer):
        for p in part.paragraphs:
            for r in list(p.runs):
                r._element.getparent().remove(r._element)


def page_field(paragraph):
    r = paragraph.add_run()
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            r._element.append(el)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._element.append(it)
    r.font.size = Pt(12)


def anchor_watermark(run, path, width_cm, x_cm, y_cm):
    """Insert a picture behind the text, absolutely positioned on the page."""
    run.add_picture(str(path), width=Cm(width_cm))
    inline = run._element.xpath(".//wp:inline")[0]
    cx, cy = inline.extent.cx, inline.extent.cy
    graphic = inline.xpath("./a:graphic")[0]
    docPr = inline.xpath("./wp:docPr")[0]
    anchor = OxmlElement("wp:anchor")
    for k, v in dict(distT="0", distB="0", distL="0", distR="0", simplePos="0", relativeHeight="0",
                     behindDoc="1", locked="0", layoutInCell="1", allowOverlap="1").items():
        anchor.set(k, v)
    sp = OxmlElement("wp:simplePos")
    sp.set("x", "0")
    sp.set("y", "0")
    anchor.append(sp)
    for tag, rel, off in (("wp:positionH", "page", x_cm), ("wp:positionV", "page", y_cm)):
        el = OxmlElement(tag)
        el.set("relativeFrom", rel)
        o = OxmlElement("wp:posOffset")
        o.text = str(int(Cm(off)))
        el.append(o)
        anchor.append(el)
    ext = OxmlElement("wp:extent")
    ext.set("cx", str(cx))
    ext.set("cy", str(cy))
    anchor.append(ext)
    ee = OxmlElement("wp:effectExtent")
    for k in ("l", "t", "r", "b"):
        ee.set(k, "0")
    anchor.append(ee)
    anchor.append(OxmlElement("wp:wrapNone"))
    anchor.append(copy.deepcopy(docPr))
    anchor.append(OxmlElement("wp:cNvGraphicFramePr"))
    anchor.append(copy.deepcopy(graphic))
    inline.getparent().replace(inline, anchor)


# ===================================================================== COVER
para(TITLE_SHORT, 18, True, align="center", after=18, line=1.15)
para("A PROJECT BASED LEARNING (PBL) REPORT", 13, True, align="center", after=12, line=1.2)
para("Submitted by", 13, True, True, align="center", after=12, line=1.2)
para(f"{M1} - {R1}", 14, True, align="center", after=8, line=1.2)
para(f"{M2} - {R2}", 14, True, align="center", after=18, line=1.2)
para("Submitted in partial fulfilment of the requirements", 13, True, True, align="center", after=6, line=1.2)
para("for the", 13, True, True, align="center", after=6, line=1.2)
para("Project-Based Learning component of Machine Learning", 13, True, True, align="center", after=14, line=1.2)
para("BACHELOR OF ENGINEERING", 15, True, align="center", after=10, line=1.2)
para("in", 11, True, True, align="center", after=10, line=1.2)
para("COMPUTER SCIENCE AND ENGINEERING", 13, True, align="center", after=12, line=1.2)
p = para("", align="center", after=10, line=1.0)
p.add_run().add_picture(str(ASSET / "p1_x71.jpeg"), width=Cm(5.6))
para("CHENNAI INSTITUTE OF TECHNOLOGY, CHENNAI", 13, True, align="center", after=2, line=1.0)
p = para("", align="center", after=6, line=1.0)
p.add_run().add_picture(str(ASSET / "p1_x72.png"), width=Cm(3.2))
para("Affiliated to Anna University, Chennai", 14, True, align="center", after=4, line=1.2)
para("(Autonomous)", 10.5, True, align="center", after=4, line=1.2)
para("OCTOBER 2026", 12, align="center", after=0, line=1.2)

# ======================================= SECTION 2: banner pages with watermark
s2 = new_section()
unlink(s2)
hp = s2.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
hp.add_run().add_picture(str(ASSET / "p2_x106.jpeg"), width=Cm(15.8))
anchor_watermark(s2.header.add_paragraph().add_run(), ASSET / "p2_x105.png", 15.0, 3.0, 12.6)

para("", after=6)
para("Vision of the Institute:", 15, True, align="left", after=10)
orange_box([(None, "To be an eminent centre for Academia, Industry and Research by imparting knowledge, relevant "
                   "practices and inculcating human values to address global challenges through novelty and sustainability.")])
para("", after=14)
para("Mission of the Institute:", 15, True, align="left", after=10)
orange_box([
    ("IM1", "To creates next generation leaders by effective teaching learning methodologies and in still Scientifics Park in them to meet the global challenges."),
    ("IM2", "To transform lives through deployment of emerging technology, novelty and sustainability."),
    ("IM3", "To inculcate human values and ethical principles to cater the societal needs."),
    ("IM4", "To contributes towards the research ecosystem by providing a suitable, effective platform for interaction between industry, academia and R & D establishments."),
    ("IM5", "To nurture incubation centres enabling structured entrepreneurship and start-ups."),
])
page_break()

para("DEPARTMENT OF", 15, True, align="center", after=0)
para("COMPUTER SCIENCE AND ENGINEERING", 15, True, align="center", after=12)
para("Vision of the Department:", 15, True, align="left", after=10)
orange_box([(None, "To Excel in the emerging areas of Computer Science and Engineering by imparting knowledge, relevant "
                   "practices and inculcating human values to transform the students as potential resources to contribute "
                   "innovatively through advanced computing in real time situations.")])
para("", after=16)
para("Mission of the Department:", 15, True, align="left", after=10)
orange_box([
    ("DM1", "To provide strong fundamentals and technical skills for Computer Science applications through effective teaching learning methodologies."),
    ("DM2", "To transform lives of the students by nurturing ethical values, creativity and novelty to become Entrepreneurs and establish start-ups."),
    ("DM3", "To habituate the students to focus on sustainable solutions to improve the quality of life and the welfare of the society."),
])
page_break()

# --- Bonafide
para("", after=4)
para("BONAFIDE CERTIFICATE", 15, True, align="center", after=14)
para(f"This is to certify that the Project–Based Learning report titled **“{TITLE}”** is a Bonafide record of work "
     f"carried out by **{m1} [{R1}]**, **{m2} [{R2}]** of the Department of Computer Science and Engineering, "
     "Chennai Institute of Technology, as part of the continuous, mentor–guided Project-Based Learning (PBL) "
     "component of the Machine Learning course during the academic year [2026–2027] under my supervision.",
     12.5, after=60, line=1.7)
sig = doc.add_table(rows=1, cols=2)
sig.alignment = WD_TABLE_ALIGNMENT.CENTER
blocks = [
    [("SIGNATURE", True), ("Dr. S. PAVITHRA, M.E., Ph.D.,", False), ("Professor and Head,", True),
     ("Dept. of Computer Science and Engineering", False), ("Chennai Institute of Technology,", False), ("Chennai – 69.", False)],
    [("SIGNATURE", True), ("Ms. S. RAJA PRIYA, M.E.,", False), ("MENTOR", True), ("Assistant Professor", True),
     ("Dept. of Computer Science and Engineering", False), ("Chennai Institute of Technology,", False), ("Chennai – 69.", False)],
]
fix_widths(sig, [7.9, 7.9])
for cell, lines in zip(sig.rows[0].cells, blocks):
    cell.text = ""
    for i, (txt, b) in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.line_spacing = 1.1
        rich(p, txt, 12.5, b)
para("", after=60)
para("Submitted for the final review held on …………………….", 12.5, align="left", after=50)
p = para("", align="left")
p.paragraph_format.tab_stops.add_tab_stop(Cm(15.8), WD_TAB_ALIGNMENT.RIGHT)
rich(p, "Internal Examiner\tExternal Examiner", 12.5, True)
page_break()

# --- Declaration
para("", after=4)
para("DECLARATION", 13, True, align="center", after=14)
para(f"I/We jointly declare that the PBL report on **“{TITLE}”** is the result of original work done by us and best "
     "of our knowledge, similar work has not been submitted to **“ANNA UNIVERSITY, CHENNAI”** for the requirement "
     "of Degree of **BACHELOR OF ENGINEERING.** This PBL report is submitted on the partial fulfilment of the "
     "requirement of the award of Degree of **COMPUTER SCIENCE AND ENGINEERING.**", 12.5, after=90, line=1.7)
para("Signature", 13, True, align="right", after=40)
para(m1, 12.5, align="right", after=40)
para(m2, 12.5, align="right", after=30)
para("Place: Chennai", 13, align="left", after=6)
para("Date: 02.10.2026", 13, align="left")
page_break()

# --- Acknowledgement
para("", after=4)
para("ACKNOWLEDGEMENT", 14, True, align="center", after=14)
for t in [
    "We wish to express our sincere gratitude to our honourable Chairman **SHRI. P. SRIRAM** for providing immense facilities at our institution.",
    "We are very proud to render our thanks to our Principal **Dr. A. RAMESH M.E., Ph.D.,** for the facilities and the encouragement given by him toward the progress and completion of our project.",
    "We would like to express special thanks and gratitude to our Dean **Dr. V. SRINIVASA RAO M.E., Ph.D.,** who has been a key source of motivation to us throughout the completion of our course and project work.",
    "We proudly render our immense gratitude to the Head of the Department **Dr. S. PAVITHRA M.E., Ph.D.,** for her effective leadership, encouragement and guidance throughout the project.",
    "We would like to extend our thanks to the Project Co-ordinator **Ms. S. RAJA PRIYA M.E.,** Department of Computer Science and Engineering, for their valuable suggestions throughout this project.",
    "We wish to acknowledge the help received from our class advisors **Ms. S. RAJA PRIYA M.E.,** and **Mr. R. RAHUL M.TECH.,** of the Department of Computer Science and Engineering for their valuable suggestions and support toward the successful completion of the project.",
]:
    para(t, 12, after=10, line=1.5)
para("", after=40)
para(f"{m1} ({R1})", 13, align="right", after=4)
para(f"{m2} ({R2})", 13, align="right")

# ======================================= SECTION 3: abstract & lists, no header
s3 = new_section()
unlink(s3)
clear_header(s3)

para("ABSTRACT", 14, True, align="center", after=10)
para(
    "Online symptom checkers are among the most widely used health tools, yet audits have found them to name the "
    "correct condition first in only about a third of cases, and most student and hobby projects in this space quote "
    "near-perfect accuracy that does not survive contact with a real patient. This project began by auditing the "
    "public disease–symptom dataset that most such projects use (4,920 records, 41 diseases, 132 binary symptoms) "
    f"and found that it contains only {E1['n_unique']} distinct records and that every one of the 41 rows in its "
    "official test file also appears in the training file, so the commonly reported 100% accuracy measures "
    "memorisation. Even under a leakage-free, profile-grouped split, standard classifiers stay near 100% when given "
    "the complete 132-symptom vector, but fall to "
    f"{pct(E2['Logistic Regression']['2'])} (logistic regression) when a patient volunteers only two symptoms, which is "
    "how real conversations begin. ANAMNESIS addresses that gap as a conversational problem rather than a "
    "classification one. A Bayesian evidence model conditions only on symptoms whose state is known, and at every turn "
    "the chatbot asks the single yes/no question with the highest expected information gain over the 41-disease "
    "posterior, stopping once one disease reaches 90% probability. Free-text replies are parsed by a lay-language "
    "symptom extractor with negation scope, warning-sign symptoms trigger an immediate emergency alert, and every "
    "answer is assembled from the knowledge base and the model’s own numbers, with no generative language model, so "
    f"nothing the chatbot says can be invented. On {E3['n_dialogues']:,} simulated consultations with held-out "
    f"patient profiles and 3% answer noise, the system reached {pct(ES['eig']['top1'])} top-1 and "
    f"{pct(ES['eig']['top3'])} top-3 accuracy after an average of {ES['eig']['mean_questions']:.1f} questions, against "
    f"{pct(ES['frequency']['top1'])} for asking the most common symptoms and {pct(ES['random']['top1'])} for random "
    f"questions, with an expected calibration error of {CAL['ece']:.3f}. The symptom extractor reached an F1 of "
    f"{NLU_F['f1']:.3f} on {E4['n_utterances']} hand-written lay-language utterances.", 12, after=10, line=1.5)
para("**Keywords:** Medical chatbot, Symptom checker, Bayesian inference, Expected information gain, Active "
     "questioning, Data leakage, Negation detection.", 12, align="left", line=1.3)
page_break()

para("TABLE OF CONTENTS", 14, True, align="center", after=12)
TOC = [
    ("CHAPTER 1: INTRODUCTION", "ch1", True, 0),
    ("1.1 Background", "s1.1", False, 0.3), ("1.2 Driving Question", "s1.2", False, 0.3),
    ("1.3 Objectives", "s1.3", False, 0.3), ("1.4 Scope and Limitations", "s1.4", False, 0.3),
    ("1.5 Novelty and Contributions", "s1.5", False, 0.3),
    ("CHAPTER 2: CONCEPT EXPLORATION", "ch2", True, 0),
    ("2.1 Related Approaches", "s2.1", False, 0.3), ("2.2 Summary Table", "s2.2", False, 0.3),
    ("2.3 What This Told Us", "s2.3", False, 0.3),
    ("CHAPTER 3: PROJECT PLANNING AND TEAM ORGANISATION", "ch3", True, 0),
    ("3.1 Weekly PBL Progress Log", "s3.1", False, 0.3), ("3.2 Requirements", "s3.2", False, 0.3),
    ("3.3 Feasibility", "s3.3", False, 0.3),
    ("CHAPTER 4: ITERATIVE DESIGN AND DEVELOPMENT", "ch4", True, 0),
    ("4.1 System Architecture", "s4.1", False, 0.3), ("4.2 Iteration 1 — Baseline", "s4.2", False, 0.3),
    ("4.3 Iteration 2 — Refinement", "s4.3", False, 0.3), ("4.4 Final Approach", "s4.4", False, 0.3),
    ("4.5 Training Procedure", "s4.5", False, 0.3),
    ("CHAPTER 5: IMPLEMENTATION", "ch5", True, 0),
    ("5.1 Module Description", "s5.1", False, 0.3), ("5.2 Key Code Snippets", "s5.2", False, 0.3),
    ("5.3 User Interface / Demo", "s5.3", False, 0.3),
    ("CHAPTER 6: RESULTS AND DISCUSSION", "ch6", True, 0),
    ("6.1 Evaluation Metrics", "s6.1", False, 0.3), ("6.2 Results Across Iterations", "s6.2", False, 0.3),
    ("6.3 Discussion", "s6.3", False, 0.3), ("6.4 Limitations", "s6.4", False, 0.3),
    ("CHAPTER 7: TEAM REFLECTION AND LEARNING OUTCOMES", "ch7", True, 0),
    ("7.1 Individual Reflections", "s7.1", False, 0.3), ("7.2 Team Learning", "s7.2", False, 0.3),
    ("7.3 Course Outcomes — Evidence Summary", "s7.3", False, 0.3),
    ("CHAPTER 8: CONCLUSION AND FUTURE SCOPE", "ch8", True, 0),
    ("8.1 Conclusion", "s8.1", False, 0.3), ("8.2 Future Scope", "s8.2", False, 0.3),
    ("REFERENCES", "refs", True, 0),
    ("APPENDIX", "app", True, 0),
    ("A.1 Full Source Code", "a1", False, 0.3), ("A.2 Complete Weekly PBL Log", "a2", False, 0.3),
    ("A.3 Self and Peer Assessment", "a3", False, 0.3), ("A.4 PBL Poster", "a4", False, 0.3),
]
for text, key, b, ind in TOC:
    toc_line(text, key, b, ind, 11.5 if not b else 12)
page_break()

FIGS = [
    ("Figure 4.1", "System Architecture Diagram", "f4.1"),
    ("Figure 4.2", "Information-Gain Questioning Loop", "f4.2"),
    ("Figure 5.1", "Consultation in Progress with Live Differential", "f5.1"),
    ("Figure 5.2", "Completed Consultation with Explanation and Triage", "f5.2"),
    ("Figure 5.3", "Warning-Sign Alert and Emergency Triage", "f5.3"),
    ("Figure 5.4", "Knowledge Question and Lay-Language Input with Negation", "f5.4"),
    ("Figure 5.5", "Phone-Width Layout", "f5.5"),
    ("Figure 6.1", "Accuracy Under Three Evaluation Protocols", "f6.1"),
    ("Figure 6.2", "Accuracy When the Patient Volunteers Only k Symptoms", "f6.2"),
    ("Figure 6.3", "Accuracy Against Number of Follow-up Questions", "f6.3"),
    ("Figure 6.4", "Reliability Diagram of Final Confidence", "f6.4"),
]
TABS = [
    ("Table 2.1", "Summary of Related Approaches", "t2.1"),
    ("Table 3.1", "Weekly PBL Progress Log", "t3.1"),
    ("Table 3.2", "Hardware and Software Requirements", "t3.2"),
    ("Table 4.1", "Model and Dialogue Configuration", "t4.1"),
    ("Table 6.1", "Leakage Audit of the Public Dataset", "t6.1"),
    ("Table 6.2", "Model Evaluation Results Across Iterations", "t6.2"),
    ("Table 6.3", "Question-Selection Policies (Early Stopping at 90%)", "t6.3"),
    ("Table 6.4", "Symptom Extraction on Lay-Language Utterances", "t6.4"),
]
para("LIST OF FIGURES", 14, True, align="center", after=12)
for a, b, k in FIGS:
    toc_line(f"{a}  {b}", k, False, 0, 11.5)
page_break()
para("LIST OF TABLES", 14, True, align="center", after=12)
for a, b, k in TABS:
    toc_line(f"{a}  {b}", k, False, 0, 11.5)
page_break()

para("LIST OF ABBREVIATIONS", 14, True, align="center", after=14)
table(["Abbreviation", "Full Form"], [
    ("ML", "Machine Learning"), ("PBL", "Project-Based Learning"), ("EIG", "Expected Information Gain"),
    ("NB", "Naive Bayes"), ("LR", "Logistic Regression"), ("RF", "Random Forest"), ("SVM", "Support Vector Machine"),
    ("k-NN", "k-Nearest Neighbours"), ("NLU", "Natural Language Understanding"), ("LLR", "Log-Likelihood Ratio"),
    ("ECE", "Expected Calibration Error"), ("TF-IDF", "Term Frequency–Inverse Document Frequency"),
    ("LLM", "Large Language Model"), ("API", "Application Programming Interface"),
    ("REST", "Representational State Transfer"), ("UI", "User Interface"), ("CSV", "Comma-Separated Values"),
], [3.6, 11.4], size=11)
page_break()

para("TEAM ROLES AND RESPONSIBILITIES", 14, True, align="center", after=12)
para("Both members jointly owned problem framing, the weekly mentor-review cycle, the literature exploration in "
     "Chapter 2 and the final report. The division below reflects primary ownership of the build; both members "
     "reviewed and tested the full system before each milestone.", 12, after=14)
table(["Team Member", "Primary Role", "Key Responsibilities"], [
    (m1, "Project Lead — ML Engine & Evaluation",
     "Dataset leakage audit and grouped split; Bayesian evidence model (model.py); expected-information-gain question "
     "selection; dialogue manager and triage rules (engine.py); simulated-patient evaluation harness and all "
     "experiments (evaluate.py); FastAPI backend; report figures."),
    (m2, "NLU, Interface & Testing",
     "Lay-language symptom lexicon and negation handling (nlu.py); the 50-utterance NLU test set; chat interface "
     "and live differential panel (HTML/CSS/JS); browser testing, screenshots and phone-width checks."),
], [3.4, 4.0, 8.4], size=11)

# ======================================= SECTION 4: main body with page numbers
s4 = new_section()
unlink(s4)
clear_header(s4)
pgNumType = OxmlElement("w:pgNumType")
pgNumType.set(qn("w:start"), "1")
s4._sectPr.append(pgNumType)
hp = s4.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
page_field(hp)

# ---------------------------------------------------------------- CHAPTER 1
heading_chapter(1, "Introduction", "ch1")
h2("1.1 Background", "s1.1")
para("For most people the first stop for a health worry is not a clinic but a phone. Symptom checkers — apps and "
     "chatbots that take a description of how someone feels and suggest what it might be and how urgently to act — "
     "are now used by millions, and in India, where the ratio of doctors to population remains low outside cities, "
     "they fill a real gap between doing nothing and travelling to a hospital. The best-known audit of these tools, "
     "however, found that across 23 commercial symptom checkers the correct diagnosis was listed first in only 34% of "
     "cases, and triage advice was appropriate in only 57% [1].")
para("A large share of academic and student work in this area is built on a single public dataset of 4,920 patient "
     "records covering 41 diseases and 132 binary symptoms [11]. Projects built on it routinely report 100% accuracy "
     "for decision trees, random forests and naive Bayes alike. Such a result should have raised suspicion. In our "
     "first iteration we audited the data and found the explanation: the 4,920 rows are copies of only "
     f"{E1['n_unique']} distinct records, and every row of the official test file is also present in the training "
     "file. The models were being tested on answers they had already been shown.")
para("Our larger observation came after removing that leakage. Given all 132 symptoms of an unseen patient, the "
     "41 diseases remain almost perfectly separable, because each disease has a distinctive signature. But no "
     "patient opens a conversation by listing 132 symptoms. They say two or three things — “fever and chills, I keep "
     "vomiting” — and the clinician’s skill lies in deciding what to ask next. This history-taking is called "
     "**anamnesis**, and it gives the project its name. ANAMNESIS treats diagnosis as a sequence of well-chosen "
     "questions rather than a single classification.")
h2("1.2 Driving Question", "s1.2")
para("The question the team set out to answer was: **“When a patient describes only a few symptoms in their own "
     "words, can a chatbot reach the right condition reliably by choosing its follow-up questions mathematically — "
     "and can it do so honestly, without inventing anything and without hiding the leakage that inflates "
     "published results?”**")
para("This narrowed into a buildable task: audit and repair the evaluation protocol for the public dataset; "
     "measure how standard classifiers behave when most symptoms are unknown; build a probabilistic model that "
     "treats unknown symptoms as unknown rather than absent; select each next question by its expected "
     "information gain; and measure, on simulated patients the model has never seen, how accuracy grows with each "
     "question asked.")
h2("1.3 Objectives", "s1.3")
for t in [
    "To audit the widely used disease–symptom dataset for duplication and train/test leakage, and to define a "
    "profile-grouped evaluation split in which no test record has a copy in training.",
    "To quantify how standard classifiers (logistic regression, random forest, SVM, naive Bayes, decision tree, "
    "k-NN) degrade when a patient reports only 1–4 symptoms, which is how real consultations begin.",
    "To design a Bayesian evidence model that conditions only on known symptoms, and an expected-information-gain "
    "policy that chooses each follow-up question to reduce diagnostic uncertainty as quickly as possible.",
    "To build a lay-language symptom extractor with negation handling, a warning-sign triage layer, and an "
    "explanation of every result in terms of the patient’s own answers.",
    "To deliver the system as a working web chatbot and evaluate it on top-1/top-3 accuracy, questions asked, "
    "calibration and robustness to answer noise.",
    "To document weekly progress and mentor feedback through the PBL cycle, and reflect on the team’s approach "
    "and division of work.",
]:
    bullet(t)
h2("1.4 Scope and Limitations", "s1.4")
para("**In scope.** A single-user web chatbot that accepts free-text symptom descriptions in English, asks yes/no "
     "follow-up questions, and returns the most consistent of 41 conditions with its probability, the evidence for "
     "and against it, a description, precautions and a triage level. It also answers “what is …” questions about "
     "the 41 conditions from its knowledge base.")
para("**Out of scope.** ANAMNESIS is not a diagnostic device and does not prescribe medicines or doses. It covers "
     "only the 41 conditions and 132 symptoms in its knowledge base and cannot recognise anything outside them. "
     "Evaluation uses simulated patients generated from held-out dataset profiles rather than real patients, and "
     "no clinician has reviewed its outputs; these limits are discussed in Section 6.4.")
h2("1.5 Novelty and Contributions", "s1.5")
para("The project makes five contributions that distinguish it from the typical symptom-prediction project built on "
     "the same data:")
for t in [
    "**A documented leakage audit.** We show that the dataset’s 4,920 rows reduce to "
    f"{E1['n_unique']} unique records and that 100% of the official test rows appear in training, and we replace "
    "the usual protocol with a profile-grouped split (Section 6.2).",
    "**Partial-evidence evaluation.** We measure accuracy when the patient volunteers only k symptoms — the "
    "condition that matters in a real conversation and that full-vector evaluation hides entirely.",
    "**Information-gain questioning.** The chatbot chooses every follow-up question by expected information gain "
    "over the disease posterior, reaching "
    f"{pct(ES['eig']['top1'])} top-1 accuracy in {ES['eig']['mean_questions']:.1f} questions on average, against "
    f"{pct(ES['frequency']['top1'])} in {ES['frequency']['mean_questions']:.1f} questions for a most-common-symptom "
    "strategy.",
    "**No generated text.** Every sentence is assembled from the knowledge base and the model’s own numbers. There is "
    "no large language model in the loop, so the chatbot cannot hallucinate a condition, a symptom or a treatment.",
    "**Visible reasoning.** A live panel shows the posterior over diseases, the remaining uncertainty in bits, the "
    "information value of each question asked, and a per-symptom likelihood-ratio explanation of the result.",
]:
    bullet(t)
page_break()

# ---------------------------------------------------------------- CHAPTER 2
heading_chapter(2, "Concept Exploration", "ch2")
h2("2.1 Related Approaches", "s2.1")
h3("2.1.1 Symptom checkers and static classifiers")
para("Semigran et al. [1] audited 23 online symptom checkers with 45 standardised patient vignettes and found the "
     "correct diagnosis listed first in 34% of evaluations and within the top 20 in 58%, with appropriate triage "
     "advice in 57%. Their study set the realistic bar for this problem. Kononenko’s survey of machine learning for "
     "medical diagnosis [2] notes that naive Bayes, despite its independence assumption, is repeatedly competitive "
     "with far more complex learners on diagnostic tasks and — importantly for a medical tool — produces "
     "explanations clinicians can follow, because each symptom contributes a separable likelihood term.")
h3("2.1.2 Data leakage in machine-learning evaluation")
para("Kaufman et al. [3] formalised leakage as information available at training time that would not be available "
     "when the model is used, and showed how easily it enters through careless splits. Kapoor and Narayanan [4] later "
     "found leakage affecting 294 published papers across 17 scientific fields, with duplicated records shared "
     "between training and test sets among the most common causes. Their taxonomy is exactly the lens that "
     "explained the 100% accuracies reported on the dataset used here.")
h3("2.1.3 Asking questions: sequential and conversational diagnosis")
para("The idea of choosing the next observation to maximise expected information goes back to Lindley [5], who "
     "defined the information an experiment provides as the expected reduction in Shannon entropy [13]. In "
     "automatic diagnosis, Wei et al. [7] framed symptom inquiry as a task-oriented dialogue learned with deep "
     "reinforcement learning and showed that an agent that asks about additional symptoms outperforms a classifier "
     "given only the patient’s self-reported symptoms (diagnosis success rate 0.65, against 0.59 accuracy for the classifier). Peng et al. [6] (REFUEL) "
     "improved such agents with reward shaping and feature rebuilding to reach higher accuracy in fewer turns on "
     "simulated patients. These agents need many thousands of training episodes and offer little insight into why "
     "a question was asked.")
h3("2.1.4 Understanding the patient’s words")
para("Patients rarely use clinical terms, and they say what they do not have as often as what they do. NegEx [8], a "
     "simple rule-based algorithm that detects negation cues and their scope in clinical text, reported 94.5% "
     "specificity and 84.5% positive predictive value — evidence that a carefully scoped rule-based approach to "
     "negation is both effective and auditable.")
h2("2.2 Summary Table", "s2.2")
table(["Ref.", "Approach / Model", "Dataset", "Reported Result"], [
    ("[1]", "23 commercial symptom checkers", "45 standardised patient vignettes", "Correct diagnosis first in 34%; appropriate triage in 57%"),
    ("[4]", "Survey of leakage in ML-based science", "294 papers, 17 fields", "Duplicates across train/test among the leading causes of inflated results"),
    ("[7]", "Deep-RL symptom inquiry dialogue", "Medical dialogue corpus (MZ)", "Success rate 0.65 vs 0.59 for a classifier on self-reported symptoms"),
    ("[6]", "REFUEL: RL with reward shaping", "Simulated patients (SymCAT)", "Higher accuracy in fewer inquiry turns than the RL baseline"),
    ("[8]", "NegEx negation detection", "Discharge summaries", "Specificity 94.5%, PPV 84.5%"),
    ("[2]", "Naive Bayes and other learners for diagnosis", "Several clinical datasets", "NB competitive with complex learners and easy to explain"),
], [1.2, 4.6, 4.4, 5.6], size=10.5)
caption("Table 2.1 — Summary of Related Approaches")
h2("2.3 What This Told Us", "s2.3")
para("Three conclusions shaped the design. First, from [3] and [4]: before trusting any accuracy figure we had to "
     "check the data for duplicates, which is what exposed the leakage in Iteration 1. Second, from [7] and [6]: the "
     "gain in symptom checking comes from asking, not from a better one-shot classifier — but reinforcement learning "
     "was unnecessary for our purpose. With an explicit probabilistic model, Lindley’s expected information gain [5] "
     "can be computed exactly for every candidate question in under a millisecond, with no training episodes, and "
     "every question can be justified in bits. Third, from [2] and [8]: in a medical setting, explainability and "
     "auditable rules outweigh small accuracy gains from opaque models. We therefore chose a naive-Bayes-style "
     "evidence model, rule-based negation, and no generative language model anywhere in the answer path.")
page_break()

# ---------------------------------------------------------------- CHAPTER 3
heading_chapter(3, "Project Planning and Team Organisation", "ch3")
h2("3.1 Weekly PBL Progress Log", "s3.1")
para("The project ran across the twelve-week PBL cycle, from the Zeroth Review on 27 June 2026, where the team "
     "presented the problem statement, to the final evaluation and report. The table summarises milestones, work "
     "completed and the mentor feedback that shaped the next iteration.")
LOG = [
    ("1–2", "Problem framing (Zeroth Review, 27 June 2026)",
     "Presented the problem statement: a medical chatbot that helps a user understand their symptoms and decide how "
     "urgently to seek care. Surveyed public symptom–disease datasets and selected the 41-disease, 132-symptom dataset.",
     "Problem approved; asked the team to state clearly that the chatbot is not a diagnostic device and to define "
     "what it must refuse to do."),
    ("3–4", "Concept exploration, baseline plan",
     "Reviewed symptom-checker audits, leakage literature and conversational diagnosis (Chapter 2); planned a "
     "classifier baseline before any chatbot work.",
     "Agreed with a measured baseline first; asked for a comparison of several classifiers rather than one."),
    ("5–6", "Iteration 1 — classifier baseline",
     "Trained six classifiers; all scored 100%. Audited the data: 4,920 rows reduce to 304 unique records and every "
     "official test row is in training. Built the profile-grouped split.",
     "Asked the team to explain the 100% before claiming it; the leakage audit was accepted as a key finding."),
    ("7–8", "Iteration 2 — partial evidence",
     "Simulated patients who volunteer only 1–4 symptoms; showed the accuracy collapse; added masked-augmentation "
     "training, which recovered part of the loss.",
     "Pointed out that real doctors ask questions rather than guess from two symptoms; suggested an interactive "
     "approach."),
    ("9–10", "Final approach — information-gain chatbot",
     "Built the Bayesian evidence model, expected-information-gain question selection, lay-language NLU with "
     "negation, warning-sign triage and the web chat interface.",
     "Asked for emergency handling for dangerous symptoms and an explanation of why each result was given."),
    ("11–12", "Evaluation, report, demo",
     "Ran the full simulated-patient evaluation, noise and calibration studies and NLU test set; captured the "
     "interface; compiled this report.",
     "Final system and evaluation accepted for the concluding review."),
]
table(["Week", "Milestone / Task", "Work Done", "Mentor Remarks"], LOG, [1.4, 3.4, 6.3, 4.7], size=10)
caption("Table 3.1 — Weekly PBL Progress Log")
h2("3.2 Requirements", "s3.2")
para("**Dataset.** The knowledge base is the public “Disease Prediction Using Machine Learning” dataset [11]: "
     f"{E1['n_rows']:,} training rows and {E1['n_official_test']} test rows, each a binary vector over "
     f"{E1['n_symptoms']} symptoms labelled with one of {E1['n_diseases']} diseases, together with companion tables "
     "of disease descriptions, four precautions per disease and a 1–7 severity weight per symptom [12]. After "
     f"de-duplication the training rows yield {E1['n_unique']} unique symptom profiles, between 5 and 10 per disease. "
     "No other labelled data was used apart from the team-written NLU test set of "
     f"{E4['n_utterances']} utterances ({E4['n_labels']} symptom labels).")
table(["Category", "Requirement"], [
    ("Processor / RAM", "Any dual-core CPU / 4 GB RAM; no GPU (developed on an Apple M-series laptop)"),
    ("Programming language", "Python 3.11+ (backend and experiments); HTML, CSS, JavaScript (interface)"),
    ("Libraries / frameworks", "NumPy, pandas, scikit-learn, Matplotlib (ML and evaluation); FastAPI, Uvicorn, "
                               "Pydantic (server); Playwright (interface testing)"),
    ("Knowledge source", "Kaggle disease–symptom dataset with description, precaution and severity tables"),
    ("Development environment", "VS Code, terminal, Chromium browser"),
    ("Version control", "Git and GitHub (repository in Appendix A.1)"),
], [4.4, 11.4], size=11)
caption("Table 3.2 — Hardware and Software Requirements")
h2("3.3 Feasibility", "s3.3")
para("The project is light on computation and heavy on experimental design, which made it feasible in the PBL "
     "window on ordinary laptops. Fitting the evidence model is a single pass of counting over 304 records; scoring "
     f"all 132 candidate questions by information gain takes about {E3['ms_per_question_eig']:.2f} ms per turn; and "
     f"the complete evaluation of {E3['n_dialogues']:,} simulated consultations under three policies, plus the noise "
     "sweep and every baseline, runs in under two minutes. The main time costs were the leakage investigation, "
     "writing the lay-language lexicon and its test set, and the interface, which were spread across the "
     "iterations in Table 3.1 so that frontend and model work could proceed in parallel.")
page_break()

# ---------------------------------------------------------------- CHAPTER 4
heading_chapter(4, "Iterative Design and Development", "ch4")
h2("4.1 System Architecture", "s4.1")
para("The system has an offline half that runs once at start-up and an online half that runs on every chat turn "
     "(Figure 4.1). Offline, the dataset is de-duplicated and the evidence model — a 41 × 132 table of the "
     "probability that a patient with each disease reports each symptom — is estimated, alongside the description, "
     "precaution and severity tables. Online, the patient’s text is converted to symptom states, the posterior over "
     "diseases is updated, and a policy decides whether to raise a warning, conclude, or ask the most informative "
     "next question.")
picture(FIG / "fig_architecture.png", 15.6)
caption("Figure 4.1 — System Architecture Diagram")
para("**Knowledge base and leakage audit:** the 4,920 rows are reduced to 304 unique profiles before anything is "
     "learned, and every evaluation splits by profile so no test record has a copy in training.")
para("**Symptom NLU:** free text is split into clauses; lay terms and misspellings are mapped to canonical symptoms "
     "by a lexicon and a character n-gram similarity match; negation cues mark symptoms as absent.")
para("**Evidence state and posterior:** every symptom is present, absent or unknown. The posterior conditions only "
     "on present and absent symptoms; unknown symptoms are marginalised out.")
para("**Policy and delivery:** warning-sign symptoms trigger an emergency alert; the conversation ends once one "
     "disease reaches 90% or questions run out; otherwise the highest-information question is asked. A FastAPI "
     "server delivers each turn to the browser, where a live panel shows the differential, evidence and triage.")
h2("4.2 Iteration 1 — Baseline", "s4.2")
para("The simplest working version followed the usual recipe for this dataset: six scikit-learn classifiers "
     "(logistic regression, random forest, RBF-kernel SVM, Bernoulli naive Bayes, decision tree and 5-nearest "
     "neighbours) trained on the 132-symptom vectors. Evaluated on the official test file and on a random 80/20 "
     "split, **every classifier scored 100%**.")
para(f"Following our mentor’s challenge to explain this, we audited the data. The {E1['n_rows']:,} training rows "
     f"contain only {E1['n_unique']} distinct records — each repeated about 16 times — and "
     f"{pct(E1['official_test_rows_seen_in_train'], 0)} of the official test rows and "
     f"{pct(E1['random_split_test_rows_seen_in_train'], 0)} of the rows in a random split also occur in training. "
     "We therefore built a **profile-grouped split**: for each disease, two unique profiles are held out entirely, "
     "over five random seeds. Under this protocol the decision tree fell to "
     f"{pct(E1['rows']['Decision Tree']['group_split_mean'])} and naive Bayes to "
     f"{pct(E1['rows']['Bernoulli Naive Bayes']['group_split_mean'])}, while logistic regression, random forest, "
     "SVM and k-NN stayed at 100% — showing that, given the full symptom vector, the diseases are genuinely "
     "separable.")
para("The real limitation surfaced when we asked what the chatbot actually receives. A patient volunteers two or "
     "three symptoms, so the other 130 inputs are not “absent” but unknown. Feeding only two reported symptoms to "
     f"the same classifiers dropped logistic regression to {pct(E2['Logistic Regression']['2'])} and random forest "
     f"to {pct(E2['Random Forest']['2'])} (Table 6.2). This motivated Iteration 2.")
h2("4.3 Iteration 2 — Refinement", "s4.3")
para("Iteration 2 attacked the partial-evidence problem in two ways. First, **masked augmentation**: each training "
     "profile was copied 20 times with a random 10–80% of its symptoms removed, teaching a logistic regression to "
     f"recognise diseases from fragments. With two reported symptoms this raised accuracy from "
     f"{pct(E2['Logistic Regression']['2'])} to {pct(E2['Masked-augmented LR']['2'])}. Second, an **evidence "
     "model** that never pretends to know unreported symptoms: a smoothed naive Bayes over present/absent/unknown "
     f"states, which reached {pct(E2['Evidence model (ours)']['2'])} from two symptoms.")
para("Both are one-shot guesses, however, and neither came close to the 100% achievable with full information. "
     "The mentor review at this stage made the decisive point: a doctor faced with “fever and vomiting” does not "
     "guess — they ask. Masked LR is the better one-shot classifier, but the evidence model has a property it lacks: "
     "for every possible question it can predict how the answer would change the diagnosis. That property is what "
     "the final approach is built on.")
h2("4.4 Final Approach", "s4.4")
para("**Evidence model.** For disease d and symptom s, θ_ds is the smoothed fraction of the disease’s profiles that "
     "include the symptom, θ_ds = (n_ds + α)/(n_d + 2α) with α = 0.05, mixed with a 3% report-slip rate ε so that a "
     "single wrong answer cannot eliminate the true disease: q_ds = θ_ds(1 − ε) + (1 − θ_ds)ε. Given evidence E, "
     "with P the set of present and A the set of absent symptoms,")
para("p(d | E) ∝ p(d) · Π_{s∈P} q_ds · Π_{s∈A} (1 − q_ds),", 12, italic=True, align="center", after=6)
para("and every unknown symptom simply drops out of the product. The prior p(d) is uniform over the 41 diseases.")
para("**Expected information gain.** For each unasked symptom s, the chatbot computes the probability of a “yes”, "
     "P(yes) = Σ_d p(d|E) q_ds, the two posteriors that would follow each answer, and the expected drop in entropy:")
para("EIG(s) = H(D | E) − [ P(yes) · H(D | E, s = yes) + P(no) · H(D | E, s = no) ]", 12, italic=True,
     align="center", after=6)
para("and asks the symptom with the largest EIG. This is Lindley’s criterion [5] applied exactly over all 132 "
     "symptoms in one vectorised NumPy operation. The conversation ends when the top disease reaches 0.90, when no "
     "question is worth at least 0.02 bits, or after eight questions. If the top disease is still below 0.50 at that "
     "point, the chatbot **abstains**: it names the three closest matches and recommends a doctor instead of "
     "naming one condition.")
picture(FIG / "fig_loop.png", 15.6)
caption("Figure 4.2 — Information-Gain Questioning Loop")
para("**Explanation.** For the final answer, each known symptom’s contribution is reported as a likelihood ratio "
     "between the top disease and its nearest rival — for example, “pain behind the eyes ×27.4” means that answer "
     "was 27 times more likely under dengue than under the runner-up. **Triage.** Ten warning-sign symptoms (chest "
     "pain, breathlessness, slurred speech, one-sided weakness, altered sensorium, coma, blood in sputum, stomach "
     "bleeding, bloody stool, acute liver failure) trigger an immediate emergency alert regardless of the diagnosis; "
     "otherwise the triage level follows the leading condition (emergency, within 24 hours, or routine).")
h2("4.5 Training Procedure", "s4.5")
para("The evidence model is fitted by counting, so there is no gradient training. Its two hyper-parameters (α and ε) "
     "and the dialogue thresholds were set on training profiles only. For every reported result, the 304 unique "
     "profiles were split by profile into 222 training and 82 test profiles (two per disease) over five seeds; each "
     "test profile was used as a **simulated patient** three times, each time volunteering two randomly chosen "
     "symptoms from its profile and answering every follow-up question truthfully except for a 3% chance of a wrong "
     f"answer. This gives {E3['n_dialogues']:,} consultations per policy.")
table(["Component", "Configuration"], [
    ("Data", f"{E1['n_unique']} unique profiles; grouped split 222 / 82 per seed; seeds 0–4; all random draws seeded"),
    ("Evidence model", "Smoothed Bernoulli likelihoods, α = 0.05; report-slip ε = 0.03; uniform prior over 41 diseases"),
    ("Question policy", "Maximum expected information gain over 132 symptoms (bits)"),
    ("Stopping rule", "Top posterior ≥ 0.90, or best EIG < 0.02 bits, or 8 questions (10 in the evaluation)"),
    ("Abstention", "Name no single disease if the final top posterior < 0.50"),
    ("Masked-augmented LR (Iter. 2)", "20 masked copies per profile, keep-rate U(0.2, 0.9); max_iter = 3000"),
    ("Baselines (Iter. 1)", "LR (max_iter 2000), RF (200 trees), SVM (RBF), Bernoulli NB, DT, k-NN (k = 5)"),
    ("Simulated patient", "2 volunteered symptoms; truthful answers with 3% flip noise (0%, 5%, 10% in sweep)"),
    ("NLU", "Lay lexicon (279 phrases) + char_wb 3–5-gram TF-IDF, cosine ≥ 0.72; clause-scoped negation"),
], [4.6, 11.2], size=10.5)
caption("Table 4.1 — Model and Dialogue Configuration")
page_break()

# ---------------------------------------------------------------- CHAPTER 5
heading_chapter(5, "Implementation", "ch5")
h2("5.1 Module Description", "s5.1")
para("**Backend — anamnesis/**", 12, align="left", after=4)
for t in [
    "**knowledge.py** — loads the dataset and its description, precaution and severity tables; cleans symptom names; "
    "de-duplicates the 4,920 rows into unique profiles; maps misspelt disease names to display names.",
    "**model.py** — the EvidenceModel: fitting the smoothed likelihood table, the partial-evidence posterior, the "
    "vectorised expected-information-gain computation and the likelihood-ratio explanation.",
    "**nlu.py** — the SymptomExtractor: clause splitting, lay-term lexicon, character n-gram fuzzy matching, "
    "negation scope, and a yes/no reply parser.",
    "**engine.py** — the dialogue manager: one Session per conversation holding the evidence vector; red-flag "
    "alerts, stopping and abstention rules, triage levels and every message template.",
    "**app/server.py** — the FastAPI application exposing /api/chat and serving the interface.",
]:
    bullet(t)
para("**Interface — app/static/**", 12, align="left", after=4, before=4)
for t in [
    "**index.html, style.css** — the chat column and the live panel (Differential, Triage, Evidence), with a "
    "responsive single-column layout below 900 px.",
    "**app.js** — renders each message type (noted symptoms, question with Yes/No buttons, warning, result card, "
    "information card, abstention) and redraws the panel from each server response.",
]:
    bullet(t)
para("**Experiments — experiments/**", 12, align="left", after=4, before=4)
bullet("**evaluate.py** — every number in Chapter 6: leakage audit, partial-evidence study, simulated-patient "
       "dialogues under three policies, noise sweep, calibration and NLU scoring; writes results.json and the "
       "figures. **nlu_testset.py** — the 50 hand-written utterances with gold labels.")
h2("5.2 Key Code Snippets", "s5.2")
para("The partial-evidence posterior. Only present and absent symptoms enter the likelihood; unknown symptoms are "
     "ignored, which is what lets the chatbot reason from two complaints:")
code('''def log_likelihood(self, state):
    pres, absn = state == PRESENT, state == ABSENT
    lt, l1t = np.log(self.theta), np.log1p(-self.theta)
    return lt[:, pres].sum(1) + l1t[:, absn].sum(1)

def posterior(self, state):
    logp = np.log(self.prior) + self.log_likelihood(state)
    logp -= logp.max()                      # numerical stability
    p = np.exp(logp)
    return p / p.sum()''')
caption("Code 5.1 — Partial-evidence Bayesian posterior (model.py)")
para("Expected information gain for all 132 candidate questions at once:")
code('''def expected_information_gain(self, post, state):
    def H(p):
        p = np.clip(p, 1e-12, 1)
        return -(p * np.log2(p)).sum(0)
    q = self.theta                                  # (diseases, symptoms)
    p_yes = post @ q                                # P(answer = yes) per symptom
    post_yes = (post[:, None] * q) / p_yes          # posterior after "yes"
    post_no = (post[:, None] * (1 - q)) / (1 - p_yes)
    eig = H(post) - (p_yes * H(post_yes) + (1 - p_yes) * H(post_no))
    eig[state != UNKNOWN] = -np.inf                 # never re-ask
    return eig''')
caption("Code 5.2 — Vectorised expected information gain (model.py)")
para("Negation scope. A cue such as “no”, “without” or “don’t have” flips every symptom found after it in the same "
     "clause; clauses are split at punctuation and at “but”, so “headache but no fever” is read correctly:")
code('''NEGATION = re.compile(r"\\b(no|not|never|without|none|denies|don't|do not|"
                      r"haven't|have not|free of|absence of)\\b")
CLAUSE_SPLIT = re.compile(r"[,.;!?]|\\bbut\\b|\\bthough\\b|\\bhowever\\b")

neg = NEGATION.search(clause)
present = not (neg and neg.start() <= match.start())''')
caption("Code 5.3 — Clause-scoped negation (nlu.py)")
para("The turn policy in the dialogue manager:")
code('''post = model.posterior(self.state)
if post.max() >= CONFIDENT and n_present >= 2 or len(self.asked) >= MAX_QUESTIONS:
    return self._conclude(model, post)        # result, or abstain if p < 0.50
eig = model.expected_information_gain(post, self.state)
j = int(np.argmax(eig))
if eig[j] < MIN_EIG_BITS:
    return self._conclude(model, post)
self.pending = j                              # ask: "Do you also have <symptom j>?"''')
caption("Code 5.4 — Conclude-or-ask policy (engine.py)")
h2("5.3 User Interface / Demo", "s5.3")
para("The interface is a single page with the conversation on the left and a live panel on the right. The patient "
     "types freely; the chatbot repeats back what it understood (“Noted: high fever; vomiting; chills; no cough”), "
     "and each follow-up question shows how many bits of information it is expected to yield, with one-tap Yes/No "
     "buttons. The panel updates after every answer: the top five diseases with their probabilities, the remaining "
     "uncertainty in bits, the evidence recorded so far (struck through when denied), each question asked with its "
     "information value, and the current triage level. All screenshots below were captured from the running "
     "application by an automated browser that typed and clicked exactly as a user would.")
picture(FIG / "ui_1_interview.png", 15.4)
caption("Figure 5.1 — Consultation in Progress with Live Differential")
para("In Figure 5.1 the patient’s first message leaves typhoid, dengue and malaria nearly tied (36%, 30%, 27%) and "
     "2.02 bits of uncertainty. The chatbot asks about loss of appetite because, of all 132 possible questions, its "
     "answer is expected to remove the most uncertainty (0.72 bits).")
picture(FIG / "ui_2_result.png", 15.4)
caption("Figure 5.2 — Completed Consultation with Explanation and Triage")
para("After two answers (Figure 5.2) dengue reaches 99% and the chatbot concludes. The result card names the "
     "runner-up, explains which answers drove the decision as likelihood ratios, and shows the description and "
     "precautions from the knowledge base; the triage panel advises seeing a doctor within 24 hours.")
picture(FIG / "ui_3_redflag.png", 15.4)
caption("Figure 5.3 — Warning-Sign Alert and Emergency Triage")
para("When the patient reports chest pain and breathlessness (Figure 5.3), a warning is raised immediately and the "
     "triage panel switches to “Seek emergency care now”, even though the differential still contains pneumonia "
     "and tuberculosis alongside heart attack — the safety rule does not wait for the model to be certain.")
picture(FIG / "ui_4_info_nlu.png", 15.4)
caption("Figure 5.4 — Knowledge Question and Lay-Language Input with Negation")
para("Figure 5.4 shows a “what is malaria?” question answered from the knowledge base, followed by everyday "
     "language — “my skin is itchy with a rash … no fever” — parsed into itching, skin rash and, through negation, "
     "no high fever.")
picture(FIG / "ui_5_mobile.png", 6.2)
caption("Figure 5.5 — Phone-Width Layout")
page_break()

# ---------------------------------------------------------------- CHAPTER 6
heading_chapter(6, "Results and Discussion", "ch6")
h2("6.1 Evaluation Metrics", "s6.1")
para("**Top-1 and top-3 accuracy** measure whether the true disease is the chatbot’s first answer or within its "
     "three most probable. Top-1 suits the single condition shown on the result card; top-3 reflects the "
     "differential a clinician would read. **Mean questions asked** measures the burden on the patient — a symptom "
     "checker that needs twenty questions will be abandoned. **Expected calibration error (ECE)** [9] measures "
     "whether a stated 90% confidence is right about 90% of the time; this matters because the chatbot uses its "
     "own confidence to decide when to stop and when to abstain. For the symptom extractor, **precision, recall "
     "and F1** are computed over (symptom, present/absent) labels, so a negated symptom read as present counts as "
     "an error, and **negation accuracy** is reported separately.")
para("All dialogue results use the profile-grouped split of Section 4.5, so every simulated patient is a symptom "
     "profile the model never saw, and every figure is the mean over five seeds. Accuracy alone would be "
     "misleading here: as Table 6.1 shows, the usual protocol reports 100% for models that have memorised the test "
     "set.")
h2("6.2 Results Across Iterations", "s6.2")
rows = []
for n, v in E1["rows"].items():
    rows.append((n, pct(v["official_test"]), pct(v["random_split"]),
                 f"{v['group_split_mean'] * 100:.1f} ± {v['group_split_std'] * 100:.1f}%"))
table(["Classifier (full 132-symptom input)", "Official Testing.csv", "Random 80/20 rows", "Profile-grouped (ours)"],
      rows, [5.6, 3.3, 3.3, 3.6], size=10.5, align_center_cols=(1, 2, 3))
caption("Table 6.1 — Leakage Audit of the Public Dataset")
para(f"Table 6.1 confirms the audit: the {E1['n_official_test']} official test rows and every row of a random split "
     "already appear in training, so the 100% column measures memory. Under a profile-grouped split the weaker "
     "learners fall, while LR, RF, SVM and k-NN remain perfect on complete vectors — the dataset is easy when every "
     "symptom is known, which is precisely why it says nothing about a conversation.")
picture(FIG / "fig_leakage.png", 15.0)
caption("Figure 6.1 — Accuracy Under Three Evaluation Protocols")
table(["Version", "Model", "Input", "Top-1", "Top-3", "Questions"], [
    ("Iteration 1 (baseline)", "Logistic regression", "2 volunteered symptoms", pct(E2["Logistic Regression"]["2"]), "—", "0"),
    ("Iteration 1 (baseline)", "Random forest", "2 volunteered symptoms", pct(E2["Random Forest"]["2"]), "—", "0"),
    ("Iteration 2 (refinement)", "Masked-augmented LR", "2 volunteered symptoms", pct(E2["Masked-augmented LR"]["2"]), "—", "0"),
    ("Iteration 2 (refinement)", "Evidence model", "2 volunteered symptoms", pct(E2["Evidence model (ours)"]["2"]),
     pct(E3["curves"]["eig"]["top3"][0]), "0"),
    ("Final approach", "Evidence model + EIG questioning", "2 symptoms + follow-up answers", pct(ES["eig"]["top1"]),
     pct(ES["eig"]["top3"]), f"{ES['eig']['mean_questions']:.1f}"),
], [3.4, 3.9, 3.5, 1.7, 1.6, 1.7], size=10, align_center_cols=(3, 4, 5))
caption("Table 6.2 — Model Evaluation Results Across Iterations")
para("Table 6.2 is the project’s central result. Starting from the same two volunteered symptoms on unseen "
     f"profiles, the best one-shot classifier reaches {pct(E2['Masked-augmented LR']['2'])}; the final chatbot, "
     f"by asking an average of {ES['eig']['mean_questions']:.1f} well-chosen questions, reaches "
     f"{pct(ES['eig']['top1'])} top-1 and {pct(ES['eig']['top3'])} top-3. Figure 6.2 shows how every one-shot model "
     "depends on how much the patient happens to volunteer.")
picture(FIG / "fig_partial.png", 15.0)
caption("Figure 6.2 — Accuracy When the Patient Volunteers Only k Symptoms")
table(["Question-selection policy", "Top-1", "Top-3", "Mean questions"], [
    ("Random unasked symptom", pct(ES["random"]["top1"]), pct(ES["random"]["top3"]), f"{ES['random']['mean_questions']:.2f}"),
    ("Most common symptom first", pct(ES["frequency"]["top1"]), pct(ES["frequency"]["top3"]), f"{ES['frequency']['mean_questions']:.2f}"),
    ("Expected information gain (ours)", pct(ES["eig"]["top1"]), pct(ES["eig"]["top3"]), f"{ES['eig']['mean_questions']:.2f}"),
], [7.0, 2.6, 2.6, 3.4], size=10.5, align_center_cols=(1, 2, 3))
caption("Table 6.3 — Question-Selection Policies (Early Stopping at 90%)")
para(f"Table 6.3 isolates the question policy, holding the model fixed. Information gain is both the most accurate "
     f"and the least burdensome: {pct(ES['eig']['top1'])} in {ES['eig']['mean_questions']:.1f} questions, against "
     f"{pct(ES['frequency']['top1'])} in {ES['frequency']['mean_questions']:.1f} for asking common symptoms first and "
     f"{pct(ES['random']['top1'])} in {ES['random']['mean_questions']:.1f} for random questions — "
     f"{ES['frequency']['mean_questions'] - ES['eig']['mean_questions']:.1f} fewer questions than the next best "
     "policy, with higher accuracy.")
picture(FIG / "fig_questions.png", 15.0)
caption("Figure 6.3 — Accuracy Against Number of Follow-up Questions")
table(["Extractor", "Precision", "Recall", "F1", "Negation accuracy"], [
    (k, f"{v['precision']:.3f}", f"{v['recall']:.3f}", f"{v['f1']:.3f}", pct(v["negation_accuracy"]))
    for k, v in E4.items() if isinstance(v, dict)
], [6.0, 2.3, 2.3, 2.0, 3.2], size=10.5, align_center_cols=(1, 2, 3, 4))
caption("Table 6.4 — Symptom Extraction on Lay-Language Utterances")
para(f"On the {E4['n_utterances']} lay-language utterances ({E4['n_labels']} labels), matching canonical names alone "
     f"finds barely a third of symptoms (recall {NLU_E['recall']:.3f}), because patients say “throwing up”, not "
     f"“vomiting”. The lay lexicon lifts F1 from {NLU_E['f1']:.3f} to {NLU_L['f1']:.3f}, and fuzzy character "
     f"n-grams, which catch misspellings such as “headach” and “diarhea”, raise it to {NLU_F['f1']:.3f} at a small "
     f"cost in precision. Negation accuracy rises from {pct(NLU_E['negation_accuracy'], 0)} to "
     f"{pct(NLU_F['negation_accuracy'], 0)}.")
h2("6.3 Discussion", "s6.3")
para("**Information gain is a long-term strategy.** Figure 6.3 shows something we did not expect: after the first "
     f"question, information gain is briefly less accurate ({pct(E3['curves']['eig']['top1'][1])}) than the "
     f"starting point ({pct(E3['curves']['eig']['top1'][0])}), while asking common symptoms improves immediately. "
     "Information gain deliberately asks the question that best splits the remaining candidates, which often "
     "spreads probability more evenly before concentrating it; the most-common-symptom policy confirms what is "
     "already likely. By the fifth question information gain overtakes it and keeps climbing to "
     f"{pct(E3['curves']['eig']['top1'][-1])} at ten questions, while the frequency policy plateaus at "
     f"{pct(E3['curves']['frequency']['top1'][-1])}. With early stopping the chatbot never stops in that dip, "
     "because the posterior is not confident there.")
para(f"**The confidence can be trusted.** The final confidence has an ECE of {CAL['ece']:.3f} (Figure 6.4), and "
     f"{pct(AB[0.9]['coverage'], 1)} of consultations end above 0.90, where the observed accuracy is "
     f"{pct(AB[0.9]['accuracy'])}. Abstaining below 0.50 withholds a single answer in only "
     f"{pct(1 - AB[0.5]['coverage'])} of cases and lifts accuracy on the rest to {pct(AB[0.5]['accuracy'])}. Below "
     "0.5 the model is under-confident rather than over-confident, which is the safe direction for a medical tool.")
picture(FIG / "fig_calibration.png", 7.6)
caption("Figure 6.4 — Reliability Diagram of Final Confidence")
para("**Robustness to wrong answers.** Patients misremember. With no answer noise the chatbot reaches "
     f"{pct(NZ['0.0']['top1'])} top-1; with 5% of answers flipped, {pct(NZ['0.05']['top1'])}; with 10%, "
     f"{pct(NZ['0.1']['top1'])} — still {pct(NZ['0.1']['top3'])} top-3 — while the mean number of questions rises "
     f"only from {NZ['0.0']['mean_questions']:.1f} to {NZ['0.1']['mean_questions']:.1f}. The 3% slip term in the "
     "model is what prevents one wrong answer from eliminating the true disease outright.")
para("**Why the evidence model and not masked LR.** Masked augmentation is the better one-shot classifier "
     f"({pct(E2['Masked-augmented LR']['2'])} against {pct(E2['Evidence model (ours)']['2'])} from two symptoms), "
     "but it cannot answer “what would I learn if I asked about joint pain?”. The evidence model can, for every "
     "symptom, in closed form, and that ability is worth far more than the one-shot gap.")
h2("6.4 Limitations", "s6.4")
for t in [
    "**Simulated patients.** All dialogue results come from patients simulated from held-out dataset profiles. Real "
    "patients describe symptoms vaguely, have more than one condition, and report symptoms outside the 132 used "
    "here; accuracy on real users will be lower.",
    "**A small, clean knowledge base.** 41 diseases and 304 unique profiles are far from clinical reality, and the "
    "dataset’s profiles are tidier than real presentations. The model inherits any errors in the source tables.",
    "**Uniform prior.** Every disease is treated as equally likely before the conversation; in practice the common "
    "cold is far more common than paralysis. Prevalence-aware priors would need epidemiological data we did not use.",
    "**No clinical validation.** No doctor has reviewed the questions, triage levels or warning-sign list, so the "
    "triage advice is a demonstration of mechanism, not a validated clinical rule.",
    "**English-only, rule-based NLU.** The extractor was tested on 50 team-written utterances; it does not handle "
    "Tamil or code-mixed input, and its negation rules can miss long-range or double negatives.",
]:
    bullet(t)
page_break()

# ---------------------------------------------------------------- CHAPTER 7
heading_chapter(7, "Team Reflection and Learning Outcomes", "ch7")
h2("7.1 Individual Reflections", "s7.1")
para(f"**{m1}:** I led the project and built the machine-learning side — the leakage audit, the evidence model, "
     "information-gain question selection, the dialogue manager and the whole evaluation harness. The biggest lesson "
     "was that a 100% accuracy figure is a question, not an answer: the most important result of the project came "
     "from refusing to accept it and counting duplicate rows. The hardest part was designing an evaluation that is "
     "fair to a conversational system — simulated patients, grouped splits, answer noise — because no standard "
     "metric existed for “how good is a chatbot at asking questions”. I also learned that deriving a formula like "
     "expected information gain by hand and vectorising it in NumPy gave me a far deeper understanding than calling "
     "a library would have.")
para(f"**{m2}:** I worked on the language and interface side — the lay-term lexicon, negation handling, the NLU "
     "test set and the chat interface with its live differential panel. Writing test utterances the way real "
     "people type showed me how little exact matching captures: canonical names alone found barely a third of the "
     "symptoms. The challenge was negation — “no fever” and “fever, no cough” look almost the same to a keyword "
     "matcher — and splitting at clause boundaries was the fix. On the interface, I learned that showing the "
     "probabilities and the information value of each question makes the system’s reasoning something a user can "
     "follow instead of a black box.")
h2("7.2 Team Learning", "s7.2")
para("Splitting the work along the boundary between model and language worked because the two halves met at a "
     "simple interface: a list of (symptom, present/absent) pairs. Either side could change without breaking the "
     "other, and the NLU could be scored on its own test set, independently of diagnosis accuracy. Building the "
     "evaluation harness early in Iteration 2 also paid off — every later design choice, from the slip rate to the "
     "stopping threshold, was decided by a number rather than an impression.")
para("Two pieces of mentor feedback changed our direction. The first, at the Iteration 1 review, was to explain the "
     "100% accuracy before claiming it, which led to the leakage audit. The second, at the Iteration 2 review, was "
     "that a doctor asks rather than guesses — the remark that turned a classification project into a "
     "conversational one. If we restarted, we would audit the dataset in week one, before training anything, and "
     "write the NLU test set before the lexicon, so that the lexicon could not be unconsciously tuned to the test.")
h2("7.3 Course Outcomes — Evidence Summary", "s7.3")
for t in [
    "**Technical / ML competency:** six supervised classifiers, a hand-derived Bayesian evidence model, expected "
    "information gain implemented from its definition, masked data augmentation and calibration analysis "
    "(Chapters 4 and 6).",
    "**Problem framing and iteration:** the progression from a leaky 100% baseline to partial-evidence models to an "
    "active-questioning chatbot, each step driven by a measured failure and dated mentor feedback (Table 3.1).",
    "**Experimental rigour:** grouped splits, five seeds, simulated patients with answer noise, ablations of the "
    "question policy and the NLU components, and a deterministic, re-runnable evaluation script.",
    "**Teamwork and division of labour:** the roles table and Section 7.1 show a clear model/language split that "
    "both members jointly reviewed and tested before each milestone.",
    "**Responsible AI:** no generated text, warning-sign triage, abstention under uncertainty, explanations for every "
    "result, and a candid statement of limitations (Section 6.4).",
]:
    bullet(t)
page_break()

# ---------------------------------------------------------------- CHAPTER 8
heading_chapter(8, "Conclusion and Future Scope", "ch8")
h2("8.1 Conclusion", "s8.1")
para("ANAMNESIS shows that the hard part of a symptom checker is not classification but conversation. On the "
     "public dataset most such projects use, we showed that the widely reported 100% accuracy comes from duplicated "
     f"records — {E1['n_rows']:,} rows hold only {E1['n_unique']} unique profiles and every official test row is in "
     "training — and that even honest classifiers fall to "
     f"{pct(E2['Logistic Regression']['2'])}–{pct(E2['Masked-augmented LR']['2'])} when a patient volunteers two "
     "symptoms. By modelling unreported symptoms as unknown and choosing each follow-up question by expected "
     f"information gain, the chatbot reached {pct(ES['eig']['top1'])} top-1 and {pct(ES['eig']['top3'])} top-3 "
     f"accuracy on unseen profiles in an average of {ES['eig']['mean_questions']:.1f} questions, with calibrated "
     f"confidence (ECE {CAL['ece']:.3f}), graceful degradation under noisy answers, and no generated text. The "
     "driving question is answered in the affirmative: well-chosen questions recover the accuracy that partial "
     "information takes away, and the system can be honest about how it got there.")
h2("8.2 Future Scope", "s8.2")
for t in [
    "Pilot the chatbot with real users under clinical supervision, and have a doctor review its questions, triage "
    "levels and warning-sign list.",
    "Replace the uniform prior with region- and season-specific prevalence (for example dengue and malaria in the "
    "monsoon), which the Bayesian design supports without retraining.",
    "Extend the knowledge base to a larger disease–symptom resource and allow symptoms with severity and duration "
    "rather than yes/no only.",
    "Add Tamil and code-mixed Tamil–English input to the NLU, and speech input for users who find typing difficult.",
    "Make the question cost-aware, trading information against how uncomfortable or hard a question is to answer, "
    "and test learned (reinforcement-learning) policies against exact information gain.",
]:
    bullet(t)
page_break()

# ---------------------------------------------------------------- REFERENCES
para("REFERENCES", 14, True, align="center", after=12)
REFS = [
    "H. L. Semigran, J. A. Linder, C. Gidengil, and A. Mehrotra, “Evaluation of symptom checkers for self diagnosis and triage: audit study,” BMJ, vol. 351, p. h3480, 2015.",
    "I. Kononenko, “Machine learning for medical diagnosis: history, state of the art and perspective,” Artificial Intelligence in Medicine, vol. 23, no. 1, pp. 89–109, 2001.",
    "S. Kaufman, S. Rosset, C. Perlich, and O. Stitelman, “Leakage in data mining: Formulation, detection, and avoidance,” ACM Trans. Knowledge Discovery from Data, vol. 6, no. 4, pp. 1–21, 2012.",
    "S. Kapoor and A. Narayanan, “Leakage and the reproducibility crisis in machine-learning-based science,” Patterns, vol. 4, no. 9, p. 100804, 2023.",
    "D. V. Lindley, “On a measure of the information provided by an experiment,” The Annals of Mathematical Statistics, vol. 27, no. 4, pp. 986–1005, 1956.",
    "Y.-S. Peng, K.-F. Tang, H.-T. Lin, and E. Chang, “REFUEL: Exploring sparse features in deep reinforcement learning for fast disease diagnosis,” in Advances in Neural Information Processing Systems 31 (NeurIPS), 2018.",
    "Z. Wei, Q. Liu, B. Peng, H. Tou, T. Chen, X. Huang, K.-F. Wong, and X. Dai, “Task-oriented dialogue system for automatic diagnosis,” in Proc. 56th Annual Meeting of the Association for Computational Linguistics (ACL), 2018, pp. 201–207.",
    "W. W. Chapman, W. Bridewell, P. Hanbury, G. F. Cooper, and B. G. Buchanan, “A simple algorithm for identifying negated findings and diseases in discharge summaries,” Journal of Biomedical Informatics, vol. 34, no. 5, pp. 301–310, 2001.",
    "C. Guo, G. Pleiss, Y. Sun, and K. Q. Weinberger, “On calibration of modern neural networks,” in Proc. 34th Int. Conf. on Machine Learning (ICML), 2017, pp. 1321–1330.",
    "F. Pedregosa et al., “Scikit-learn: Machine learning in Python,” Journal of Machine Learning Research, vol. 12, pp. 2825–2830, 2011.",
    "Kaushil268, “Disease Prediction Using Machine Learning,” Kaggle dataset, 2020. [Online]. Available: https://www.kaggle.com/datasets/kaushil268/disease-prediction-using-machine-learning",
    "itachi9604, “healthcare-chatbot: symptom description, precaution and severity tables,” GitHub repository. [Online]. Available: https://github.com/itachi9604/healthcare-chatbot",
    "C. E. Shannon, “A mathematical theory of communication,” Bell System Technical Journal, vol. 27, no. 3, pp. 379–423, 1948.",
    "S. Ramírez, “FastAPI documentation.” [Online]. Available: https://fastapi.tiangolo.com/ (accessed Sept. 2026).",
]
for i, r in enumerate(REFS, 1):
    p = para("", align="justify", after=5, line=1.3)
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.9)
    rich(p, f"[{i}] {r}", 11.5)
page_break()

# ---------------------------------------------------------------- APPENDIX
para("APPENDIX", 14, True, align="center", after=12)
h2("A.1 Full source code:")
para(REPO, 12, align="left", color="1F4E9A")
para("The repository contains the chatbot (anamnesis/, app/), the evaluation harness (experiments/), the scripts "
     "that drew every figure and captured every screenshot (scripts/), the raw results (results/results.json) and "
     "instructions to reproduce each table with a single command.", 12)
h2("A.2 Complete weekly PBL log and mentor sign-offs:")
para("The full week-by-week log is given in Table 3.1 (Chapter 3); it spans the twelve-week PBL cycle beginning at "
     "the Zeroth Review on 27 June 2026.", 12)
h2("A.3 Self and Peer Assessment")
table(["Team Member", "Self-Rated Contribution (%)", "Peer-Rated Contribution (%)", "Remarks"], [
    (m1, "60%", "60%", "Led the project; leakage audit, evidence model, information-gain dialogue, evaluation and backend"),
    (m2, "40%", "40%", "Lay-language NLU and negation, NLU test set, chat interface and interface testing"),
], [3.6, 3.0, 3.0, 6.2], size=10.5, align_center_cols=(1, 2))
page_break()
h2("A.4 PBL Poster")
if (ROOT / "poster" / "poster_rotated.png").exists():
    picture(ROOT / "poster" / "poster_rotated.png", 15.2)

out = ROOT / "report" / "PBL_Report_Medical_Chatbot_ANAMNESIS.docx"
doc.save(out)
print("saved", out)
